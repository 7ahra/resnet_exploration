#!/usr/bin/env python3
"""BCCE 2026 deck: Chemistry faculty beliefs about doctoral education."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SPEECH = os.path.join(HERE, "assets", "speech_bubbles_transparent.png")
MAP = os.path.join(HERE, "assets", "theme_boundary_map.png")

# ---------------- Design system ----------------
INK   = RGBColor(0x24,0x29,0x28)
MUTE  = RGBColor(0x74,0x72,0x6A)
FAINT = RGBColor(0xB6,0xB2,0xA7)
BG    = RGBColor(0xFC,0xFB,0xF7)
CARD  = RGBColor(0xF3,0xF0,0xE8)
LINE  = RGBColor(0xE4,0xE0,0xD5)
DARK  = RGBColor(0x1B,0x1E,0x1D)
DARKMUTE = RGBColor(0xA9,0xA5,0x9A)
WHITE = RGBColor(0xFF,0xFF,0xFF)

T1 = RGBColor(0x45,0x6A,0x73)   # teal
T2 = RGBColor(0xA3,0x77,0x30)   # gold
T3 = RGBColor(0x5E,0x7B,0x4F)   # green
T4 = RGBColor(0x9D,0x64,0x73)   # mauve
THEMES = [T1, T2, T3, T4]

SERIF = "Georgia"
SANS  = "Calibri"

EMU_IN = 914400
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5
BLANK = prs.slide_layouts[6]

# ---------------- helpers ----------------
def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    # send to back
    sp = r._element; sp.getparent().remove(sp); s.shapes._spTree.insert(2, sp)
    return s

def _set_spacing(run, spc):
    rPr = run._r.get_or_add_rPr(); rPr.set('spc', str(spc))

def txt(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        line_spacing=1.0, space_after=0):
    """runs: list of paragraphs; each paragraph = list of (text, dict) tuples."""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('left','right','top','bottom'):
        setattr(tf, 'margin_'+m, 0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing: p.line_spacing = line_spacing
        p.space_before = 0; p.space_after = Pt(space_after)
        for (text, st) in para:
            r = p.add_run(); r.text = text
            f = r.font
            f.name = st.get('font', SANS)
            f.size = Pt(st.get('size', 18))
            f.bold = st.get('bold', False)
            f.italic = st.get('italic', False)
            f.color.rgb = st.get('color', INK)
            if 'spc' in st: _set_spacing(r, st['spc'])
    return tb

def kicker(slide, text, color, l=0.92, t=0.62):
    txt(slide, l, t, 11.5, 0.4, [[(text.upper(), dict(font=SANS, size=12.5, bold=True,
        color=color, spc=220))]])

def footer(slide, n, dark=False):
    c = DARKMUTE if dark else FAINT
    lc = RGBColor(0x33,0x37,0x36) if dark else LINE
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.92), Inches(7.02),
                                    Inches(SW-0.92), Inches(7.02))
    ln.line.color.rgb = lc; ln.line.width = Pt(0.75); ln.shadow.inherit = False
    txt(slide, 0.92, 7.08, 8, 0.3, [[("Faculty beliefs about doctoral education", dict(size=10, color=c)),
        ("   ·   BCCE 2026", dict(size=10, color=c, italic=True))]])
    txt(slide, SW-2.0, 7.08, 1.08, 0.3, [[(f"{n:02d}", dict(size=10, color=c, bold=True))]],
        align=PP_ALIGN.RIGHT)

def rrect(slide, l, t, w, h, fill=None, line=None, line_w=1.0, radius=0.08, shadow=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    try: sh.adjustments[0] = radius
    except Exception: pass
    return sh

def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def chip(slide, l, t, text, color, filled=False, size=11.5, w=None, pad=0.16, h=0.36):
    tw = (len(text)*0.083*size/11.5 + pad*2) if w is None else w
    sh = rrect(slide, l, t, tw, h, fill=(color if filled else None),
               line=(None if filled else color), line_w=1.25, radius=0.5)
    tc = WHITE if filled else color
    txt(slide, l, t, tw, h, [[(text, dict(size=size, bold=True, color=tc, font=SANS))]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return tw

def axis(slide, l, t, w, color, left, right, caption=None, dark=False):
    y = t
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(l), Inches(y),
                                      Inches(l+w), Inches(y))
    conn.line.color.rgb = color; conn.line.width = Pt(2.0); conn.shadow.inherit = False
    lnEl = conn.line._get_or_add_ln()
    for tag in ('a:headEnd', 'a:tailEnd'):
        e = lnEl.makeelement(qn(tag), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        lnEl.append(e)
    lc = INK if not dark else RGBColor(0xEC,0xE8,0xDE)
    txt(slide, l-0.15, y+0.08, w*0.5, 0.4, [[(left, dict(size=14, color=lc, font=SANS, bold=True))]],
        align=PP_ALIGN.LEFT)
    txt(slide, l+w*0.5+0.15, y+0.08, w*0.5, 0.4, [[(right, dict(size=14, color=lc, font=SANS, bold=True))]],
        align=PP_ALIGN.RIGHT)
    if caption:
        cc = MUTE if not dark else DARKMUTE
        txt(slide, l, y+0.52, w, 0.5, [[(caption, dict(size=12.5, italic=True, color=cc))]],
            align=PP_ALIGN.CENTER)

def bigquote(slide, l, t, w, color, quote, attrib, size=25, mark=True, dark=False):
    if mark:
        txt(slide, l-0.06, t-0.55, 1.2, 1.2, [[("“", dict(font=SERIF, size=90, color=color, bold=True))]])
    qc = INK if not dark else RGBColor(0xF0,0xEC,0xE2)
    txt(slide, l+0.62, t+0.05, w-0.62, 2.6,
        [[(quote, dict(font=SERIF, size=size, italic=True, color=qc))]], line_spacing=1.12)
    ac = MUTE if not dark else DARKMUTE
    return

# =================================================================
# SLIDE 1 — TITLE
# =================================================================
s = slide()
rect(s, 0, 0, 0.28, SH, T1)  # thin spine accent (will layer other colors)
rect(s, 0, 0, 0.28, SH*0.25, T1)
rect(s, 0, SH*0.25, 0.28, SH*0.25, T2)
rect(s, 0, SH*0.50, 0.28, SH*0.25, T3)
rect(s, 0, SH*0.75, 0.28, SH*0.25, T4)
txt(s, 0.95, 0.72, 11.5, 0.4, [[("BCCE 2026", dict(size=13, bold=True, color=T1, spc=240)),
    ("   ·   CHEMISTRY EDUCATION RESEARCH", dict(size=13, bold=True, color=MUTE, spc=240))]])
txt(s, 0.92, 1.55, 11.6, 2.6,
    [[("What chemistry faculty believe", dict(font=SERIF, size=52, color=INK))],
     [("about doctoral education", dict(font=SERIF, size=52, color=INK))]], line_spacing=1.02)
txt(s, 0.95, 3.35, 11.5, 0.6,
    [[("A phenomenographic and thematic characterization", dict(font=SERIF, size=21, italic=True, color=T1))]])
# authors
txt(s, 0.95, 4.35, 11.5, 0.5,
    [[("Zoe Rahimi Pirkoohi", dict(size=18, bold=True, color=INK)),
      ("      Jordan Harshman", dict(size=18, color=INK))]])
txt(s, 0.95, 4.82, 11.5, 0.4,
    [[("University of Iowa", dict(size=14.5, color=MUTE))]])
# speech bubble figure across the bottom right
s.shapes.add_picture(SPEECH, Inches(6.9), Inches(4.55), height=Inches(2.7))
txt(s, 0.95, 7.06, 11.5, 0.35,
    [[("NSF CAREER #2142873 & #2602955  ·  Jordan Harshman, PI", dict(size=10.5, color=FAINT))]])

# =================================================================
# SLIDE 2 — MOTIVATION
# =================================================================
s = slide()
kicker(s, "The enterprise", T1)
txt(s, 0.92, 1.15, 11.5, 1.6,
    [[("Doctoral education produces the chemists", dict(font=SERIF, size=36, color=INK))],
     [("who staff the discipline.", dict(font=SERIF, size=36, color=INK))]], line_spacing=1.05)
# three stat blocks
stats = [("~2,900", "chemistry PhDs\nawarded each year", T1),
         ("~196", "U.S. doctoral\nprograms", T2),
         ("Decades", "the same structure,\nlargely unchanged", T3)]
x = 0.92; bw = 3.62; gap = 0.28
for (big, small, col) in stats:
    rrect(s, x, 3.15, bw, 1.75, fill=CARD, radius=0.06)
    rect(s, x, 3.15, 0.10, 1.75, col)
    txt(s, x+0.35, 3.36, bw-0.5, 0.9, [[(big, dict(font=SERIF, size=40, bold=True, color=col))]])
    lines = small.split("\n")
    txt(s, x+0.37, 4.12, bw-0.5, 0.7,
        [[(l, dict(size=14, color=MUTE))] for l in lines], line_spacing=1.05)
    x += bw + gap
txt(s, 0.92, 5.35, 11.5, 1.2,
    [[("The same record documents its strains — ", dict(size=18, color=INK)),
      ("competing responsibilities, no standard assessment of outcomes, uneven implementation, and inconsistent mentorship.",
       dict(size=18, color=INK))]], line_spacing=1.25)
txt(s, 0.92, 6.45, 11.5, 0.4, [[("Collini et al., 2025  ·  Donkor & Harshman, 2023", dict(size=12, italic=True, color=FAINT))]])
footer(s, 2)

# =================================================================
# SLIDE 3 — THE GAP
# =================================================================
s = slide()
kicker(s, "The gap", T3)
txt(s, 0.92, 1.12, 11.5, 1.5,
    [[("We know what faculty believe about ", dict(font=SERIF, size=31, color=INK)),
      ("teaching a class.", dict(font=SERIF, size=31, italic=True, color=T3))],
     [("We don’t know what they believe about ", dict(font=SERIF, size=31, color=INK)),
      ("advising a PhD.", dict(font=SERIF, size=31, italic=True, color=T1))]], line_spacing=1.12)
# two columns
cy = 3.35; ch = 2.75
rrect(s, 0.92, cy, 5.55, ch, fill=CARD, radius=0.05)
rect(s, 0.92, cy, 0.10, ch, T3)
txt(s, 1.25, cy+0.28, 5.0, 0.5, [[("CLASSROOM TEACHING", dict(size=13, bold=True, color=T3, spc=180))]])
txt(s, 1.25, cy+0.85, 5.0, 1.8,
    [[("Faculty beliefs about teaching and learning are well characterized — and predict what happens in the classroom.",
       dict(size=16.5, color=INK))],
     [("Gibbons · Lund & Stains · Popova · Mack & Towns",
       dict(size=12.5, italic=True, color=MUTE))]], line_spacing=1.22, space_after=10)

rrect(s, 6.85, cy, 5.55, ch, fill=CARD, radius=0.05)
rect(s, 6.85, cy, 0.10, ch, T1)
txt(s, 7.18, cy+0.28, 5.0, 0.5, [[("RESEARCH ADVISING", dict(size=13, bold=True, color=T1, spc=180))]])
txt(s, 7.18, cy+0.85, 5.0, 1.8,
    [[("The actual work of the doctorate is a ", dict(size=16.5, color=INK)),
      ("different practice", dict(size=16.5, bold=True, color=T1)),
      (" — and beliefs are tied to the practice they concern. Never characterized.",
       dict(size=16.5, color=INK))],
     [("Kagan, 1992", dict(size=12.5, italic=True, color=MUTE))]], line_spacing=1.22, space_after=10)
# punchline
txt(s, 0.92, 6.35, 11.5, 0.6,
    [[("Donkor et al. (2024) set faculty beliefs deliberately aside — a proper investigation ", dict(size=15, color=INK)),
      ("“warranted its own study.”", dict(size=15, italic=True, bold=True, color=T1)),
      ("   This is that study.", dict(size=15, color=INK))]])
footer(s, 3)

# =================================================================
# SLIDE 4 — RESEARCH QUESTION
# =================================================================
s = slide(bg=T1)
txt(s, 1.2, 1.0, 10.9, 0.5, [[("RESEARCH QUESTION", dict(size=14, bold=True, color=RGBColor(0xCF,0xDD,0xE1), spc=280))]],
    align=PP_ALIGN.CENTER)
txt(s, 1.0, 2.45, 11.3, 2.4,
    [[("What do chemistry faculty believe", dict(font=SERIF, size=42, color=WHITE))],
     [("about doctoral education?", dict(font=SERIF, size=42, color=WHITE))]],
    align=PP_ALIGN.CENTER, line_spacing=1.08)
# four areas as a row
areas = ["What the doctorate is for", "How students learn", "Who develops which skill", "The structural conditions"]
x = 1.15; bw = 2.63; gap = 0.22
for a in areas:
    txt(s, x, 5.35, bw, 0.9, [[(a, dict(size=14.5, color=RGBColor(0xE6,0xEE,0xF0), font=SANS))]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, line_spacing=1.1)
    x += bw + gap
# little divider dots
for i in range(3):
    dx = 1.15 + bw*(i+1) + gap*i + gap/2 - 0.03
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(dx), Inches(5.55), Inches(0.07), Inches(0.07))
    d.fill.solid(); d.fill.fore_color.rgb = RGBColor(0x9C,0xB8,0xBF); d.line.fill.background(); d.shadow.inherit=False

# =================================================================
# SLIDE 5 — WHY IT MATTERS (TCSR)
# =================================================================
s = slide()
kicker(s, "Why beliefs matter  ·  Teacher-Centered Systemic Reform", T4)
txt(s, 0.92, 1.12, 11.6, 1.5,
    [[("Reforms keep changing the structures of schooling", dict(font=SERIF, size=29, color=INK))],
     [("without changing what happens inside them.", dict(font=SERIF, size=29, color=INK))]], line_spacing=1.1)
# four elements feeding practice
labels = ["General context\nof reform", "Teacher's\npersonal profile", "Structural &\ncultural context", "Teacher thinking\n(knowledge + beliefs)"]
x = 0.92; bw = 2.72; gap = 0.2; ytop = 3.2; bh = 1.35
for i, lb in enumerate(labels):
    hot = (i == 3)
    col = T4 if hot else CARD
    rrect(s, x, ytop, bw, bh, fill=(T4 if hot else CARD), radius=0.07)
    tc = WHITE if hot else INK
    lines = lb.split("\n")
    txt(s, x+0.15, ytop, bw-0.3, bh, [[(l, dict(size=15, bold=hot, color=tc))] for l in lines],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    x += bw + gap
txt(s, 0.92, 4.75, 11.5, 0.5, [[("↓   all four shape whether a reform reaches ", dict(size=15, color=MUTE)),
    ("practice", dict(size=15, italic=True, bold=True, color=INK)),
    (" — the ultimate determinant of its success", dict(size=15, color=MUTE))]], align=PP_ALIGN.CENTER)
# takeaway pill
rrect(s, 2.6, 5.55, 8.13, 0.95, fill=RGBColor(0xF4,0xEC,0xEE), radius=0.5)
txt(s, 2.9, 5.55, 7.5, 0.95, [[("Beliefs are self-perpetuating — they can ", dict(size=17, color=INK)),
    ("absorb a reform without yielding to it.", dict(size=17, bold=True, color=T4))]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 5)

# =================================================================
# SLIDE 6 — WHAT IS A BELIEF (decision rule / Fig 2)
# =================================================================
s = slide()
kicker(s, "What counts as a belief", T2)
txt(s, 0.92, 1.12, 11.6, 1.1,
    [[("A ", dict(font=SERIF, size=27, color=INK)),
      ("proposition a person holds true", dict(font=SERIF, size=27, italic=True, color=T2)),
      (" and uses to interpret experience and guide action.", dict(font=SERIF, size=27, color=INK))]],
    line_spacing=1.1)
txt(s, 0.92, 2.5, 11.5, 0.4, [[("Rokeach, 1968  ·  Green, 1971  ·  Nespor, 1987  ·  Pajares, 1992", dict(size=12.5, italic=True, color=FAINT))]])
txt(s, 0.92, 3.15, 11.5, 0.4, [[("A unit of talk counts as a belief only if it passes all four tests:", dict(size=15.5, color=MUTE))]])
# four gates
gates = [("1", "Propositional", "a proposition\nheld to be true"),
         ("2", "Evaluative", "guides interpretation\n& decisions"),
         ("3", "Personal", "a personal stance,\nnot group consensus"),
         ("4", "Resistant", "holds up against\nits own counter-evidence")]
x = 0.92; bw = 2.55; gap = 0.42; ytop = 3.75; bh = 1.85
for i, (num, title, sub) in enumerate(gates):
    rrect(s, x, ytop, bw, bh, fill=WHITE, line=T2, line_w=1.5, radius=0.08)
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+bw/2-0.28), Inches(ytop+0.22), Inches(0.56), Inches(0.56))
    d.fill.solid(); d.fill.fore_color.rgb = T2; d.line.fill.background(); d.shadow.inherit=False
    txt(s, x+bw/2-0.28, ytop+0.22, 0.56, 0.56, [[(num, dict(size=20, bold=True, color=WHITE))]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+0.1, ytop+0.92, bw-0.2, 0.4, [[(title, dict(size=16, bold=True, color=INK))]], align=PP_ALIGN.CENTER)
    txt(s, x+0.1, ytop+1.3, bw-0.2, 0.6, [[(l, dict(size=12.5, color=MUTE))] for l in sub.split("\n")],
        align=PP_ALIGN.CENTER, line_spacing=1.05)
    if i < 3:
        ar = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x+bw+0.06), Inches(ytop+bh/2-0.16), Inches(0.3), Inches(0.32))
        ar.fill.solid(); ar.fill.fore_color.rgb = FAINT; ar.line.fill.background(); ar.shadow.inherit=False
    x += bw + gap
# result pill
rrect(s, 4.55, 5.95, 4.23, 0.62, fill=T2, radius=0.5)
txt(s, 4.55, 5.95, 4.23, 0.62, [[("= a belief", dict(size=17, bold=True, color=WHITE))]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 6)

# =================================================================
# SLIDE 7 — METHODS
# =================================================================
s = slide()
kicker(s, "Methods", T1)
txt(s, 0.92, 1.12, 11.5, 0.9,
    [[("An interpretive design built to capture the ", dict(font=SERIF, size=27, color=INK)),
      ("range", dict(font=SERIF, size=27, italic=True, color=T1)),
      (" of faculty belief", dict(font=SERIF, size=27, color=INK))]], line_spacing=1.05)
# left: what we did
rows = [("20 faculty", "across U.S. doctoral programs, research-active, currently advising"),
        ("Interviews + card sort", "~60-min semi-structured (modified Teacher Beliefs Interview); the sort as a second elicitation"),
        ("Phenomenography + thematic analysis", "the range of ways faculty understand, plus keyness — not frequency — as the reporting rule")]
y = 2.45
for (h, sub) in rows:
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.98), Inches(y+0.06), Inches(0.16), Inches(0.16))
    d.fill.solid(); d.fill.fore_color.rgb = T1; d.line.fill.background(); d.shadow.inherit=False
    txt(s, 1.35, y-0.05, 10.9, 0.4, [[(h, dict(size=18.5, bold=True, color=INK))]])
    txt(s, 1.35, y+0.36, 10.7, 0.6, [[(sub, dict(size=14, color=MUTE))]], line_spacing=1.12)
    y += 1.02
# pipeline
py = 5.5; ph = 0.95
rrect(s, 0.92, py, 11.5, ph, fill=CARD, radius=0.35)
pipe = [("990", "sub-codes"), ("15", "codes"), ("4", "themes")]
centers = [3.0, 6.67, 10.34]
for (big, small), cx in zip(pipe, centers):
    txt(s, cx-1.7, py, 3.4, ph, [[(big+"  ", dict(font=SERIF, size=32, bold=True, color=T1)),
        (small, dict(size=16, color=INK))]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for mx in (4.835, 8.505):
    txt(s, mx-0.35, py, 0.7, ph, [[("→", dict(size=26, color=FAINT))]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.92, py+ph+0.15, 11.5, 0.4,
    [[("No new codes appeared after interview 7 — the result held from 10 participants to all 20.",
       dict(size=13, italic=True, color=MUTE))]], align=PP_ALIGN.CENTER)
footer(s, 7)

# =================================================================
# SLIDE 8 — THE SAMPLE
# =================================================================
s = slide()
kicker(s, "The sample  ·  N = 20", T3)
txt(s, 0.92, 1.12, 11.5, 0.9,
    [[("Recruited for ", dict(font=SERIF, size=27, color=INK)),
      ("maximum variation", dict(font=SERIF, size=27, italic=True, color=T3)),
      (" across profile and subfield", dict(font=SERIF, size=27, color=INK))]])
# proportion bars
bars = [("Gender", 13, 20, "13 men  ·  7 women"),
        ("Race / ethnicity", 13, 20, "13 White  ·  7 non-White"),
        ("Rank", 15, 20, "15 full professors"),
        ("Institution", 14, 20, "14 at R1 institutions"),
        ("Doctoral training", 18, 20, "18 trained in the U.S.")]
y = 2.55; bw = 5.4; lx = 0.92; bh = 0.34
for (lab, num, tot, note) in bars:
    txt(s, lx, y-0.03, 2.3, 0.4, [[(lab, dict(size=13.5, bold=True, color=INK))]])
    bx = lx + 2.35
    rrect(s, bx, y, bw, bh, fill=RGBColor(0xEC,0xE9,0xE0), radius=0.5)
    rrect(s, bx, y, bw*num/tot, bh, fill=T3, radius=0.5)
    txt(s, bx+bw+0.15, y-0.03, 3.2, 0.4, [[(note, dict(size=12.5, color=MUTE))]])
    y += 0.62
# subdiscipline chips
txt(s, 0.92, 5.85, 11.5, 0.4, [[("SUBDISCIPLINES", dict(size=11.5, bold=True, color=T3, spc=200))]])
subs = [("Inorganic", 5), ("Chem. education", 2), ("Materials", 2), ("Physical", 2),
        ("Biophysical", 2), ("Organic", 2), ("+5 singletons", 1)]
x = 0.92; cy = 6.28
for (nm, n) in subs:
    label = f"{nm}" + (f" ×{n}" if n > 1 else "")
    w = chip(s, x, cy, label, T3, filled=False, size=11.5)
    x += w + 0.16
txt(s, 0.92, 6.85, 11.5, 0.3, [[("Analytical generalization & transferability — not population estimates. Skews senior/male; see limitations.",
    dict(size=10.5, italic=True, color=FAINT))]])
footer(s, 8)

# =================================================================
# SLIDE 9 — THE HEADLINE FINDING
# =================================================================
s = slide(bg=DARK)
txt(s, 1.2, 1.05, 10.9, 0.5, [[("THE FINDING", dict(size=14, bold=True, color=DARKMUTE, spc=300))]], align=PP_ALIGN.CENTER)
txt(s, 0.9, 2.2, 11.5, 2.2,
    [[("Faculty ", dict(font=SERIF, size=40, color=WHITE)),
      ("agree", dict(font=SERIF, size=40, italic=True, color=RGBColor(0x9C,0xC0,0xC9))),
      (" on the propositions.", dict(font=SERIF, size=40, color=WHITE))],
     [("They ", dict(font=SERIF, size=40, color=WHITE)),
      ("divide", dict(font=SERIF, size=40, italic=True, color=RGBColor(0xD6,0xA1,0xB0))),
      (" on where the boundaries fall.", dict(font=SERIF, size=40, color=WHITE))]],
    align=PP_ALIGN.CENTER, line_spacing=1.12)
txt(s, 2.0, 5.15, 9.3, 1.2,
    [[("“What differs is not which beliefs faculty hold, but where they draw the lines those beliefs require.”",
       dict(font=SERIF, size=18, italic=True, color=RGBColor(0xD8,0xD4,0xC9)))]],
    align=PP_ALIGN.CENTER, line_spacing=1.2)

# =================================================================
# SLIDES 10-13 — FOUR THEMES
# =================================================================
theme_data = [
    dict(col=T1, n=1, name="What the doctorate is for",
         prop=["The doctorate develops a self-directed scientist", "in service of the student’s own goals."],
         codes=["IC","UN","SG","FE"],
         quote="When they start bringing ideas forward without necessarily being prompted.",
         attrib="P2  ·  on what independence looks like",
         left="Answering existing questions", right="Asking new questions",
         cap="The realistic goal vs. the higher bar for the degree"),
    dict(col=T2, n=2, name="How students learn",
         prop=["Students learn by doing scaffolded research —", "and learning shows in performance, not grades."],
         codes=["LE","LO","AV"],
         quote="I do. We do. You do — spread out over X number of years.",
         attrib="P5  ·  on how research is learned",
         left="Unsupported struggle helps", right="Struggle should be scaffolded",
         cap="Same formative deprivation, opposite lessons — P13 vs. P7"),
    dict(col=T3, n=3, name="Who develops which skill  (largest theme)",
         prop=["Development is a shared responsibility in a two-way", "relationship — and who develops which skill varies."],
         codes=["MB","FA","OE","RS","SD","MP"],
         quote="I go from the leader to the follower, and they start as the follower and end up as the leader.",
         attrib="P6  ·  on how responsibility transfers",
         left="Faculty’s to teach", right="Student’s to build",
         cap="The same skill — e.g. creativity — is one faculty’s to teach (P5, P17), another’s for the student (P9, P11)"),
    dict(col=T4, n=4, name="The structural conditions",
         prop=["Structure shapes advising — and the gaps it leaves", "are the program’s to fix, not the student’s."],
         codes=["PS","GA"],
         quote="Those professional skills often get kind of forgotten.",
         attrib="P2  ·  on the program’s gap",
         left="The evolved system works", right="Curricula & funding need change",
         cap="How much of the structure deserves changing — P14 vs. P18"),
]
for idx, d in enumerate(theme_data):
    col = d["col"]
    s = slide()
    rect(s, 0, 0, 0.22, SH, col)
    txt(s, 0.92, 0.6, 8, 0.4, [[(f"THEME {d['n']} OF 4", dict(size=12.5, bold=True, color=col, spc=240)),
        ("   ·   "+d["name"].upper(), dict(size=12.5, bold=True, color=MUTE, spc=160))]])
    # code chips top-right
    cx = SW - 0.92
    for code in reversed(d["codes"]):
        w = 0.62
        cx -= (w + 0.12)
    x = cx
    for code in d["codes"]:
        wch = chip(s, x, 0.55, code, col, filled=True, size=12, w=0.62, h=0.4)
        x += 0.62 + 0.12
    # proposition
    txt(s, 0.92, 1.25, 11.5, 1.4,
        [[(line, dict(font=SERIF, size=30, color=INK))] for line in d["prop"]], line_spacing=1.08)
    # quote card
    rrect(s, 0.92, 2.95, 11.5, 1.95, fill=CARD, radius=0.05)
    rect(s, 0.92, 2.95, 0.10, 1.95, col)
    bigquote(s, 1.35, 3.25, 10.6, col, d["quote"], d["attrib"], size=24)
    txt(s, 1.97, 4.45, 10.0, 0.4, [[("— " + d["attrib"], dict(size=14, bold=True, color=col))]])
    # boundary axis
    txt(s, 0.92, 5.2, 11.5, 0.4, [[("WHERE FACULTY DIVIDE", dict(size=12, bold=True, color=col, spc=200))]])
    axis(s, 2.4, 5.95, 8.5, col, d["left"], d["right"], caption=d["cap"])
    footer(s, 10+idx)

# =================================================================
# SLIDE 14 — THE FULL OUTCOME SPACE (Fig 3, dark)
# =================================================================
s = slide(bg=DARK)
txt(s, 0.7, 0.55, 12, 0.4, [[("THE FULL OUTCOME SPACE", dict(size=12.5, bold=True, color=DARKMUTE, spc=260))]])
# figure on left
img_h = 6.2
pic = s.shapes.add_picture(MAP, Inches(0.6), Inches(1.05), height=Inches(img_h))
# right-side annotation
rx = 6.9
txt(s, rx, 1.4, 5.9, 2.0,
    [[("4 shared themes.", dict(font=SERIF, size=30, color=WHITE))],
     [("15 codes.", dict(font=SERIF, size=30, color=WHITE))],
     [("An axis of division in every one.", dict(font=SERIF, size=30, italic=True, color=RGBColor(0xD8,0xD4,0xC9)))]],
    line_spacing=1.12)
txt(s, rx, 4.3, 5.9, 2.0,
    [[("Every faculty member lands somewhere on each axis. The agreement is real — and so is the disagreement one level down.",
       dict(size=16, color=DARKMUTE))]], line_spacing=1.3)
txt(s, rx, 6.35, 5.9, 0.4, [[("Each axis names the two poles of variation; nothing is plotted on it.", dict(size=11.5, italic=True, color=RGBColor(0x6E,0x6B,0x63)))]])

# =================================================================
# SLIDE 15 — WHAT THIS ADDS
# =================================================================
s = slide()
kicker(s, "What this adds", T1)
txt(s, 0.92, 1.12, 11.6, 0.9,
    [[("A belief baseline for changing doctoral education", dict(font=SERIF, size=29, color=INK))]])
cards = [
    (T1, "Content for the model", "Gives the Teacher-Centered Systemic Reform model the doctoral-education content it lacked — the shared ground a reform can assume, and the fault lines it must engage."),
    (T2, "Nests with prior work", "The goals record (what the degree is for) and the challenge record (what goes wrong) now sit beside what faculty believe. Three answers, one enterprise — they nest, not compete."),
    (T3, "Conviction vs. practice", "13 of 20 describe stable convictions but implementation refined by experience. Supports aimed at implementation may fit faculty experience better than efforts aimed at core beliefs."),
]
x = 0.92; bw = 3.72; gap = 0.27; ytop = 2.35; bh = 3.9
for (col, h, body) in cards:
    rrect(s, x, ytop, bw, bh, fill=CARD, radius=0.05)
    rect(s, x, ytop, bw, 0.12, col)
    txt(s, x+0.32, ytop+0.42, bw-0.6, 0.9, [[(h, dict(font=SERIF, size=20, bold=True, color=col))]], line_spacing=1.0)
    txt(s, x+0.32, ytop+1.35, bw-0.62, 2.4, [[(body, dict(size=14.5, color=INK))]], line_spacing=1.28)
    x += bw + gap
footer(s, 15)

# =================================================================
# SLIDE 16 — IMPLICATIONS + THANK YOU
# =================================================================
s = slide(bg=T1)
txt(s, 0.95, 0.7, 11.5, 0.5, [[("IMPLICATIONS", dict(size=13, bold=True, color=RGBColor(0xCF,0xDD,0xE1), spc=260))]])
txt(s, 0.92, 1.35, 11.5, 1.4,
    [[("Two levels of the finding,", dict(font=SERIF, size=32, color=WHITE))],
     [("two lessons for programs.", dict(font=SERIF, size=32, color=WHITE))]], line_spacing=1.05)
# two blocks
rrect(s, 0.92, 3.0, 5.55, 2.3, fill=RGBColor(0x3A,0x5B,0x63), radius=0.06)
txt(s, 1.25, 3.25, 5.0, 0.5, [[("Shared propositions", dict(size=15, bold=True, color=RGBColor(0xBF,0xD4,0xD9), spc=100))]])
txt(s, 1.25, 3.75, 5.0, 1.5, [[("Common ground a reform argument can take as given — arguing from where faculty already stand.",
    dict(size=16.5, color=WHITE))]], line_spacing=1.25)
rrect(s, 6.85, 3.0, 5.55, 2.3, fill=WHITE, radius=0.06)
txt(s, 7.18, 3.25, 5.0, 0.5, [[("Boundary disagreements", dict(size=15, bold=True, color=T1, spc=100))]])
txt(s, 7.18, 3.75, 5.0, 1.5, [[("The harder, more useful target — where colleagues who sound aligned quietly diverge. Surface them, don’t assume them.",
    dict(size=16.5, color=INK))]], line_spacing=1.25)
# thank you / ack
txt(s, 0.92, 5.65, 11.5, 0.6, [[("Thank you.", dict(font=SERIF, size=26, italic=True, color=WHITE))]])
txt(s, 0.92, 6.45, 11.5, 0.4,
    [[("Zoe Rahimi Pirkoohi & Jordan Harshman  ·  University of Iowa", dict(size=13, color=RGBColor(0xDCE6E8 if False else 0xDC,0xE6,0xE8)))]])
txt(s, 0.92, 6.85, 11.5, 0.35,
    [[("Supported by NSF CAREER #2142873 & #2602955  ·  Jordan Harshman, PI", dict(size=10.5, color=RGBColor(0xA9,0xC1,0xC7)))]])

out = os.path.join(HERE, "BCCE2026_faculty_beliefs.pptx")
prs.save(out)
print("saved", out, "with", len(prs.slides._sldIdLst), "slides")
