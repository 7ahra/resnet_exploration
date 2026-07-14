#!/usr/bin/env python3
"""BCCE 2026 deck v2 — restrained academic redesign.
White base, near-black text, sparing flamingo (#F5628C) + Iowa gold (#FFCD00).
Constantia (display/serif) + Corbel (labels/sans)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SPEECH = os.path.join(HERE, "assets", "speech_bubbles_gray.png")
MAP = os.path.join(HERE, "assets", "boundary_map_v2.png")

# palette
BG    = RGBColor(0xFC,0xFB,0xF9)
INK   = RGBColor(0x1A,0x17,0x12)
GRAY  = RGBColor(0x6B,0x66,0x60)
FAINT = RGBColor(0xA8,0xA2,0x99)
LINE  = RGBColor(0xE4,0xE0,0xD6)
PINK  = RGBColor(0xF5,0x62,0x8C)
PINKD = RGBColor(0xDA,0x37,0x69)   # deeper flamingo for smaller emphasis on white
GOLD  = RGBColor(0xFF,0xCD,0x00)
BLACK = RGBColor(0x12,0x11,0x10)
WHITE = RGBColor(0xFF,0xFF,0xFF)

SERIF = "Constantia"
SANS  = "Corbel"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5
BLANK = prs.slide_layouts[6]
LM = 1.0  # left margin

def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background(); r.shadow.inherit = False
    sp = r._element; sp.getparent().remove(sp); s.shapes._spTree.insert(2, sp)
    return s

def _spc(run, v): run._r.get_or_add_rPr().set('spc', str(v))

def txt(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        line_spacing=1.0, space_after=0):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('left','right','top','bottom'): setattr(tf, 'margin_'+m, 0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing: p.line_spacing = line_spacing
        p.space_before = 0; p.space_after = Pt(space_after)
        for (text, st) in para:
            r = p.add_run(); r.text = text; f = r.font
            f.name = st.get('font', SANS); f.size = Pt(st.get('size', 18))
            f.bold = st.get('bold', False); f.italic = st.get('italic', False)
            f.color.rgb = st.get('color', INK)
            if 'spc' in st: _spc(r, st['spc'])
    return tb

def hrule(slide, x, y, w, color=LINE, weight=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x+w), Inches(y))
    c.line.color.rgb = color; c.line.width = Pt(weight); c.shadow.inherit = False; return c

def vrule(slide, x, y, h, color=LINE, weight=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y+h))
    c.line.color.rgb = color; c.line.width = Pt(weight); c.shadow.inherit = False; return c

def vbar(slide, x, y, h, w=0.055, color=GOLD):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background(); r.shadow.inherit = False; return r

def axis(slide, l, y, w, left, right, caption=None, color=PINK, labelcolor=INK, capcolor=GRAY):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(l), Inches(y), Inches(l+w), Inches(y))
    c.line.color.rgb = color; c.line.width = Pt(1.75); c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    for tag in ('a:headEnd','a:tailEnd'):
        ln.append(ln.makeelement(qn(tag), {'type':'triangle','w':'med','len':'med'}))
    txt(slide, l-0.1, y+0.1, w*0.55, 0.4, [[(left, dict(size=14, color=labelcolor))]])
    txt(slide, l+w*0.45+0.1, y+0.1, w*0.55, 0.4, [[(right, dict(size=14, color=labelcolor))]], align=PP_ALIGN.RIGHT)
    if caption:
        txt(slide, l, y+0.52, w, 0.4, [[(caption, dict(size=12.5, italic=True, color=capcolor))]])

def footer(slide, n, dark=False):
    if n is None: return
    c1 = FAINT if not dark else RGBColor(0x77,0x72,0x69)
    lc = LINE if not dark else RGBColor(0x33,0x30,0x2B)
    hrule(slide, LM, 7.06, SW-2*LM, color=lc, weight=0.75)
    txt(slide, LM, 7.13, 8, 0.3,
        [[("What chemistry faculty believe about doctoral education", dict(size=9.5, italic=True, color=c1, font=SERIF))]])
    # gold tick + number
    t = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(SW-LM-0.42), Inches(7.17), Inches(0.08), Inches(0.08))
    t.fill.solid(); t.fill.fore_color.rgb = GOLD; t.line.fill.background(); t.shadow.inherit=False
    txt(slide, SW-LM-0.3, 7.11, 0.3, 0.3, [[(f"{n}", dict(size=10, color=c1, bold=True))]], align=PP_ALIGN.RIGHT)

def microlabel(slide, text, x=LM, y=0.66, color=GRAY):
    txt(slide, x, y, 10, 0.3, [[(text.upper(), dict(size=10.5, color=color, spc=200))]])

# =============================================================== S1 TITLE
s = slide()
txt(s, LM, 0.6, 6, 0.35, [[("BCCE 2026", dict(size=12, color=GRAY, spc=260))]])
txt(s, SW-LM-6, 0.6, 6, 0.35, [[("UNIVERSITY OF IOWA", dict(size=12, color=INK, spc=260, bold=True))]], align=PP_ALIGN.RIGHT)
hrule(s, LM, 1.02, SW-2*LM, color=GOLD, weight=2.5)
vbar(s, LM, 2.62, 1.72, w=0.06, color=GOLD)
txt(s, LM+0.32, 2.5, 11.4, 2.1,
    [[("What chemistry faculty ", dict(font=SERIF, size=46, color=INK)),
      ("believe", dict(font=SERIF, size=46, color=INK, italic=True))],
     [("about doctoral education", dict(font=SERIF, size=46, color=INK))]], line_spacing=1.04)
txt(s, LM+0.32, 4.45, 11, 0.5,
    [[("A phenomenographic and thematic characterization", dict(font=SERIF, size=20, italic=True, color=GRAY))]])
txt(s, LM+0.32, 5.45, 11, 0.4,
    [[("Zoe Rahimi Pirkoohi", dict(size=16.5, color=INK, bold=True)),
      ("    ·    Jordan Harshman", dict(size=16.5, color=GRAY))]])
txt(s, LM+0.32, 5.86, 11, 0.35, [[("University of Iowa", dict(size=13, color=GRAY))]])
s.shapes.add_picture(SPEECH, Inches(8.55), Inches(5.0), height=Inches(1.95))
txt(s, LM+0.32, 7.12, 11, 0.35,
    [[("Supported by NSF CAREER #2142873 & #2602955  ·  Jordan Harshman, PI", dict(size=10, color=FAINT))]])

# =============================================================== S2 ENTERPRISE
s = slide()
txt(s, LM, 1.15, 11.2, 1.4,
    [[("Doctoral education produces the chemists", dict(font=SERIF, size=33, color=INK))],
     [("who staff the discipline.", dict(font=SERIF, size=33, color=INK))]], line_spacing=1.06)
nums = [("~2,900", "chemistry PhDs\nawarded each year"),
        ("~196", "U.S. doctoral\nprograms"),
        ("decades", "one structure,\nlargely unchanged")]
colw = (SW-2*LM)/3
for i,(big,small) in enumerate(nums):
    x = LM + i*colw
    if i>0: vrule(s, x-0.15, 3.35, 1.5, color=LINE, weight=1.0)
    txt(s, x, 3.35, colw-0.3, 0.9, [[(big, dict(font=SERIF, size=44, color=INK))]])
    txt(s, x, 4.28, colw-0.3, 0.8, [[(l, dict(size=13.5, color=GRAY))] for l in small.split("\n")], line_spacing=1.08)
txt(s, LM, 5.55, 11.0, 1.0,
    [[("The same record documents its strains — competing responsibilities, no standard "
       "assessment of outcomes, uneven implementation, inconsistent mentorship.",
       dict(size=16, color=GRAY, font=SERIF))]], line_spacing=1.3)
txt(s, LM, 6.5, 11, 0.3, [[("Donkor & Harshman, 2023  ·  Collini et al., 2025", dict(size=11, italic=True, color=FAINT))]])
footer(s, 2)

# =============================================================== S3 THE GAP
s = slide()
txt(s, LM, 1.1, 11.4, 1.7,
    [[("We know what faculty believe about ", dict(font=SERIF, size=29, color=INK)),
      ("teaching a class.", dict(font=SERIF, size=29, italic=True, color=GRAY))],
     [("We don’t know what they believe about ", dict(font=SERIF, size=29, color=INK)),
      ("advising a PhD.", dict(font=SERIF, size=29, color=PINKD))]], line_spacing=1.14)
vrule(s, SW/2, 3.5, 2.0, color=LINE, weight=1.0)
txt(s, LM, 3.5, SW/2-LM-0.4, 2.0,
    [[("Classroom teaching", dict(size=15.5, bold=True, color=INK))],
     [("Beliefs are well characterized — and predict what happens in the classroom.",
       dict(size=15, color=GRAY, font=SERIF))],
     [("Gibbons · Lund & Stains · Popova · Mack & Towns", dict(size=11.5, italic=True, color=FAINT))]],
    line_spacing=1.24, space_after=7)
txt(s, SW/2+0.4, 3.5, SW/2-LM-0.4, 2.0,
    [[("Research advising", dict(size=15.5, bold=True, color=INK))],
     [("A different practice — and beliefs are tied to the practice they concern. Never characterized.",
       dict(size=15, color=GRAY, font=SERIF))],
     [("Kagan, 1992", dict(size=11.5, italic=True, color=FAINT))]], line_spacing=1.24, space_after=7)
txt(s, LM, 6.1, 11.3, 0.6,
    [[("Donkor et al. (2024) set faculty beliefs aside — a proper investigation ", dict(size=15, color=INK, font=SERIF)),
      ("“warranted its own study.”", dict(size=15, italic=True, color=GRAY, font=SERIF)),
      ("  This is that study.", dict(size=15, color=PINKD, font=SERIF, bold=True))]])
footer(s, 3)

# =============================================================== S4 RESEARCH QUESTION
s = slide()
hrule(s, LM, 1.72, 0.62, color=PINK, weight=2.5)
microlabel(s, "Research question", y=1.95, color=GRAY)
txt(s, LM, 2.75, 10.8, 2.6,
    [[("What do chemistry faculty ", dict(font=SERIF, size=40, color=INK)),
      ("believe", dict(font=SERIF, size=40, color=INK, italic=True))],
     [("about doctoral education?", dict(font=SERIF, size=40, color=INK))]], line_spacing=1.08)
txt(s, LM, 5.55, 11.3, 0.5,
    [[("what it’s for   ·   how students learn   ·   who develops which skill   ·   the structure they work in",
       dict(size=14.5, color=GRAY))]])
footer(s, 4)

# =============================================================== S5 TCSR
s = slide()
microlabel(s, "Why beliefs matter · Teacher-Centered Systemic Reform")
txt(s, LM, 1.15, 11.4, 1.5,
    [[("Reforms change the structures of schooling", dict(font=SERIF, size=28, color=INK))],
     [("without changing what happens inside them.", dict(font=SERIF, size=28, color=INK))]], line_spacing=1.1)
els = ["general\ncontext of reform", "the teacher’s\npersonal profile", "structural &\ncultural context", "teacher thinking\n→ beliefs"]
colw = (SW-2*LM)/4
for i,e in enumerate(els):
    x = LM+i*colw
    if i>0: vrule(s, x-0.1, 3.55, 1.15, color=LINE, weight=1.0)
    hot = (i==3)
    lines = e.split("\n")
    runs=[]
    for j,l in enumerate(lines):
        runs.append([(l, dict(size=15.5, color=(PINKD if hot else INK), bold=hot, font=SANS))])
    txt(s, x, 3.6, colw-0.25, 1.1, runs, line_spacing=1.12)
txt(s, LM, 4.95, 11, 0.4, [[("All four shape whether a reform reaches ", dict(size=14, color=GRAY)),
    ("practice", dict(size=14, italic=True, color=INK, font=SERIF)),
    (" — the model’s ultimate test.", dict(size=14, color=GRAY))]])
hrule(s, LM, 5.65, SW-2*LM, color=LINE, weight=0.75)
txt(s, LM, 5.85, 11.3, 0.9,
    [[("Beliefs are self-perpetuating — they can ", dict(font=SERIF, size=21, color=INK)),
      ("absorb a reform without yielding to it.", dict(font=SERIF, size=21, color=PINKD, italic=True))]], line_spacing=1.15)
footer(s, 5)

# =============================================================== S6 WHAT IS A BELIEF
s = slide()
microlabel(s, "What counts as a belief")
txt(s, LM, 1.12, 11.5, 1.3,
    [[("A ", dict(font=SERIF, size=25, color=INK)),
      ("proposition a person holds true", dict(font=SERIF, size=25, color=PINKD)),
      (" and uses to interpret", dict(font=SERIF, size=25, color=INK))],
     [("experience and guide action.", dict(font=SERIF, size=25, color=INK))]], line_spacing=1.12)
txt(s, LM, 2.6, 11, 0.3, [[("Rokeach, 1968 · Green, 1971 · Nespor, 1987 · Pajares, 1992", dict(size=11.5, italic=True, color=FAINT))]])
txt(s, LM, 3.2, 11, 0.35, [[("A unit of talk counts as a belief only when it is all four:", dict(size=14.5, color=GRAY))]])
rules = [("1","Propositional","a proposition the person holds to be true"),
         ("2","Evaluative","it guides how they interpret experience and decide"),
         ("3","Personal","a personal stance, not a claim needing group consensus"),
         ("4","Resistant","it holds up against their own counter-evidence")]
y=3.75
for (n,term,gloss) in rules:
    txt(s, LM, y, 0.5, 0.4, [[(n, dict(font=SERIF, size=20, color=PINK, bold=True))]])
    txt(s, LM+0.55, y+0.03, 3.2, 0.4, [[(term, dict(size=16.5, bold=True, color=INK))]])
    txt(s, LM+3.7, y+0.05, 8.0, 0.4, [[(gloss, dict(size=15, color=GRAY, font=SERIF))]])
    y+=0.72
footer(s, 6)

# =============================================================== S7 METHODS
s = slide()
microlabel(s, "Methods")
txt(s, LM, 1.12, 11.5, 0.9,
    [[("An interpretive design built to capture the ", dict(font=SERIF, size=26, color=INK)),
      ("range", dict(font=SERIF, size=26, color=PINKD, italic=True)),
      (" of faculty belief", dict(font=SERIF, size=26, color=INK))]])
rows=[("Twenty faculty","across U.S. doctoral programs — research-active, currently advising a PhD student."),
      ("Interviews + a card sort","~60-minute semi-structured interviews (a modified Teacher Beliefs Interview); the sort as a second elicitation."),
      ("Phenomenography + thematic analysis","the range of ways faculty understand doctoral education, with keyness — not frequency — as the reporting rule.")]
y=2.55
for (h,g) in rows:
    txt(s, LM, y, 11.3, 0.4, [[(h, dict(size=18, bold=True, color=INK))]])
    txt(s, LM, y+0.38, 11.3, 0.5, [[(g, dict(size=14, color=GRAY, font=SERIF))]], line_spacing=1.12)
    y+=1.0
hrule(s, LM, 5.7, SW-2*LM, color=LINE, weight=0.75)
pipe=[("990","sub-codes"),("15","codes"),("4","themes")]
x=LM+0.1
for i,(big,small) in enumerate(pipe):
    txt(s, x, 5.9, 2.2, 0.7, [[(big+" ", dict(font=SERIF, size=30, color=INK, bold=True)),
        (small, dict(size=15, color=GRAY))]], anchor=MSO_ANCHOR.MIDDLE)
    if i<2:
        txt(s, x+2.15, 5.9, 0.7, 0.7, [[("→", dict(size=22, color=PINK))]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    x+=2.85
txt(s, 9.4, 5.9, 3.0, 0.7, [[("no new codes after interview 7", dict(size=13, italic=True, color=GRAY, font=SERIF))]],
    anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
footer(s, 7)

# =============================================================== S8 SAMPLE
s = slide()
microlabel(s, "The sample · N = 20")
txt(s, LM, 1.12, 11.5, 0.9,
    [[("Recruited for ", dict(font=SERIF, size=26, color=INK)),
      ("maximum variation", dict(font=SERIF, size=26, color=PINKD, italic=True)),
      (" across profile and subfield", dict(font=SERIF, size=26, color=INK))]])
bars=[("Gender",13,"13 men · 7 women"),("Race / ethnicity",13,"13 White · 7 non-White"),
      ("Rank",15,"15 full professors"),("Institution",14,"14 at R1 institutions"),
      ("Doctoral training",18,"18 trained in the U.S.")]
y=2.5; bw=5.0; bx=LM+2.55; bh=0.26
for (lab,num,note) in bars:
    txt(s, LM, y-0.05, 2.4, 0.4, [[(lab, dict(size=13, color=INK))]])
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx), Inches(y), Inches(bw), Inches(bh))
    r.fill.solid(); r.fill.fore_color.rgb=LINE; r.line.fill.background(); r.shadow.inherit=False
    r2=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx), Inches(y), Inches(bw*num/20), Inches(bh))
    r2.fill.solid(); r2.fill.fore_color.rgb=INK; r2.line.fill.background(); r2.shadow.inherit=False
    txt(s, bx+bw+0.2, y-0.06, 3.4, 0.4, [[(note, dict(size=12.5, color=GRAY))]])
    y+=0.56
txt(s, LM, 5.55, 11.4, 0.4, [[("SUBDISCIPLINES", dict(size=10.5, color=GRAY, spc=180))]])
txt(s, LM, 5.9, 11.4, 0.4,
    [[("Inorganic (5)  ·  Chemistry education (2)  ·  Materials (2)  ·  Physical (2)  ·  Biophysical (2)  ·  Organic (2)  ·  + five singletons",
       dict(size=13, color=INK, font=SERIF))]])
txt(s, LM, 6.5, 11.4, 0.3,
    [[("Analytical generalization and transferability, not population estimates; the sample skews senior and male — see limitations.",
       dict(size=11, italic=True, color=FAINT))]])
footer(s, 8)

# =============================================================== S9 THE FINDING (black)
s = slide(bg=BLACK)
txt(s, LM, 1.35, 8, 0.35, [[("THE FINDING", dict(size=12, color=RGBColor(0x9A,0x94,0x8B), spc=280))]])
txt(s, LM, 2.55, 11.5, 2.4,
    [[("Faculty ", dict(font=SERIF, size=40, color=WHITE)),
      ("agree", dict(font=SERIF, size=40, color=GOLD, italic=True)),
      (" on the propositions.", dict(font=SERIF, size=40, color=WHITE))],
     [("They ", dict(font=SERIF, size=40, color=WHITE)),
      ("divide", dict(font=SERIF, size=40, color=PINK, italic=True)),
      (" on where the boundaries fall.", dict(font=SERIF, size=40, color=WHITE))]], line_spacing=1.14)
txt(s, LM, 5.5, 10.8, 1.0,
    [[("“What differs is not which beliefs faculty hold, but where they draw the lines those beliefs require.”",
       dict(font=SERIF, size=17, italic=True, color=RGBColor(0xC9,0xC3,0xB9)))]], line_spacing=1.25)

# =============================================================== S10-13 THEMES
themes=[
 dict(n=1,name="What the doctorate is for",codes="IC · UN · SG · FE",
   prop=[("The doctorate develops a ",INK),("self-directed scientist",PINKD),("",INK)],
   prop2="in service of the student’s own goals.",
   quote="When they start bringing ideas forward without necessarily being prompted.",
   who="P2, on what independence looks like",
   left="answering existing questions", right="asking new questions",
   cap="The realistic goal versus the higher bar for the degree."),
 dict(n=2,name="How students learn",codes="LE · LO · AV",
   prop=[("Students learn by doing ",INK),("scaffolded research",PINKD),("",INK)],
   prop2="— and learning shows in performance, not grades.",
   quote="I do. We do. You do — spread out over X number of years.",
   who="P5, on how research is learned",
   left="unsupported struggle helps", right="struggle should be scaffolded",
   cap="Same formative deprivation, opposite lessons — P13 vs. P7."),
 dict(n=3,name="Who develops which skill — the largest theme",codes="MB · FA · OE · RS · SD · MP",
   prop=[("Development is a ",INK),("shared responsibility",PINKD),("",INK)],
   prop2="in a two-way relationship — and who develops which skill varies.",
   quote="I go from the leader to the follower, and they start as the follower and end up as the leader.",
   who="P6, on how responsibility transfers",
   left="faculty’s to teach", right="student’s to build",
   cap="The same skill — creativity — is one faculty’s to teach (P5, P17), another’s for the student (P9, P11)."),
 dict(n=4,name="The structural conditions",codes="PS · GA",
   prop=[("Structure shapes advising — and the ",INK),("gaps",PINKD),("",INK)],
   prop2="it leaves are the program’s to fix, not the student’s.",
   quote="Those professional skills often get kind of forgotten.",
   who="P2, on the program’s gap",
   left="the evolved system works", right="curricula & funding need change",
   cap="How much of the structure deserves changing — P14 vs. P18."),
]
for i,d in enumerate(themes):
    s = slide()
    microlabel(s, f"Theme {d['n']} / 4    ·    {d['name']}")
    prop_runs=[(t, dict(font=SERIF, size=29, color=c)) for (t,c) in d["prop"] if t]
    txt(s, LM, 1.15, 11.4, 1.5, [prop_runs, [(d["prop2"], dict(font=SERIF, size=29, color=INK))]], line_spacing=1.1)
    # quote with thin pink left rule
    vbar(s, LM, 3.15, 1.25, w=0.045, color=PINK)
    txt(s, LM+0.3, 3.15, 10.9, 1.3, [[(d["quote"], dict(font=SERIF, size=22, italic=True, color=INK))]], line_spacing=1.16)
    txt(s, LM+0.3, 4.35, 10.9, 0.35, [[("— "+d["who"].split(",")[0], dict(size=13, bold=True, color=INK)),
        (","+d["who"].split(",",1)[1], dict(size=13, color=GRAY))]])
    # boundary
    txt(s, LM, 5.15, 8, 0.3, [[("WHERE THEY DIVIDE", dict(size=10.5, color=GRAY, spc=200))]])
    axis(s, LM+0.1, 5.75, 10.6, d["left"], d["right"], caption=d["cap"])
    txt(s, LM, 6.62, 11, 0.3, [[(d["codes"], dict(size=10.5, color=FAINT, spc=60))]])
    footer(s, 10+i)

# =============================================================== S14 OUTCOME MAP
s = slide()
s.shapes.add_picture(MAP, Inches(0.62), Inches(0.7), height=Inches(6.35))
RX=7.15
microlabel(s, "The full outcome space", x=RX, y=1.15)
txt(s, RX, 1.7, 5.4, 2.4,
    [[("Four shared themes.", dict(font=SERIF, size=27, color=INK))],
     [("Fifteen codes.", dict(font=SERIF, size=27, color=INK))],
     [("An ", dict(font=SERIF, size=27, color=INK)),
      ("axis of division", dict(font=SERIF, size=27, color=PINKD, italic=True))],
     [("in every one.", dict(font=SERIF, size=27, color=INK))]], line_spacing=1.12)
txt(s, RX, 4.55, 5.4, 1.8,
    [[("The agreement is real — and so is the disagreement one level down. Every faculty member "
       "lands somewhere on each axis.", dict(size=15, color=GRAY, font=SERIF))]], line_spacing=1.3)

# =============================================================== S15 WHAT THIS ADDS
s = slide()
microlabel(s, "What this adds")
txt(s, LM, 1.12, 11.5, 0.9,
    [[("A ", dict(font=SERIF, size=27, color=INK)),
      ("belief baseline", dict(font=SERIF, size=27, color=PINKD)),
      (" for changing doctoral education", dict(font=SERIF, size=27, color=INK))]])
pts=[("1","Content for the model","Gives the Teacher-Centered Systemic Reform model the doctoral-education content it lacked — the shared ground a reform can assume, and the fault lines it must engage."),
     ("2","Nests with prior work","The goals record and the challenge record now sit beside what faculty believe. Three answers to one enterprise — they nest, not compete."),
     ("3","Conviction vs. practice","Thirteen of twenty hold stable convictions but refine implementation with experience — so supports aimed at implementation may fit faculty best.")]
y=2.6
for (n,h,g) in pts:
    txt(s, LM, y, 0.55, 0.5, [[(n, dict(font=SERIF, size=22, color=PINK, bold=True))]])
    txt(s, LM+0.6, y+0.02, 11.0, 0.4, [[(h, dict(size=18, bold=True, color=INK))]])
    txt(s, LM+0.6, y+0.44, 10.8, 0.7, [[(g, dict(size=14.5, color=GRAY, font=SERIF))]], line_spacing=1.2)
    y+=1.35
footer(s, 15)

# =============================================================== S16 IMPLICATIONS + THANKS
s = slide()
microlabel(s, "Implications")
txt(s, LM, 1.12, 11.5, 0.8, [[("Two levels, two lessons.", dict(font=SERIF, size=30, color=INK))]])
vrule(s, SW/2, 2.5, 1.9, color=LINE, weight=1.0)
txt(s, LM, 2.5, SW/2-LM-0.4, 2.0,
    [[("SHARED PROPOSITIONS", dict(size=11, color=GRAY, spc=160))],
     [("Common ground a reform can take as given — arguing from where faculty already stand.",
       dict(size=16.5, color=INK, font=SERIF))]], line_spacing=1.24, space_after=8)
txt(s, SW/2+0.4, 2.5, SW/2-LM-0.4, 2.0,
    [[("BOUNDARY DISAGREEMENTS", dict(size=11, color=GRAY, spc=160))],
     [("The harder, more useful target — where colleagues who sound aligned ", dict(size=16.5, color=INK, font=SERIF)),
      ("quietly diverge.", dict(size=16.5, color=PINKD, font=SERIF, italic=True)),
      (" Surface them; don’t assume them.", dict(size=16.5, color=INK, font=SERIF))]], line_spacing=1.24, space_after=8)
hrule(s, LM, 4.95, SW-2*LM, color=LINE, weight=0.75)
vbar(s, LM, 5.35, 0.85, w=0.06, color=GOLD)
txt(s, LM+0.3, 5.35, 10, 0.9, [[("Thank you.", dict(font=SERIF, size=30, italic=True, color=INK))]])
txt(s, LM+0.3, 6.35, 11.3, 0.35,
    [[("Zoe Rahimi Pirkoohi  ·  Jordan Harshman  ·  University of Iowa", dict(size=13, color=GRAY))]])
txt(s, LM+0.3, 6.75, 11.3, 0.3,
    [[("Supported by NSF CAREER #2142873 & #2602955  ·  Jordan Harshman, PI", dict(size=10, color=FAINT))]])

out = os.path.join(HERE, "BCCE2026_faculty_beliefs.pptx")
prs.save(out)
print("saved", out, "with", len(prs.slides._sldIdLst), "slides")
