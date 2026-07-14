#!/usr/bin/env python3
"""BCCE 2026 deck v4 — assertion-evidence STORYTELLING in the lab house style.
Every title is a claim; titles chain into the argument; concepts are built
before use; the 'agree but divide' finding is dramatized with opposing quotes."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(HERE, "assets")
BGTEX = os.path.join(A, "bg_texture.png")
MAP = os.path.join(A, "boundary_map_v3.png")

GREEN = RGBColor(0x7D,0x92,0x63); GREEND = RGBColor(0x5C,0x6E,0x47)
DARK  = RGBColor(0x2C,0x2C,0x2C); GRAY = RGBColor(0x59,0x59,0x59)
FAINT = RGBColor(0x8C,0x89,0x83); WHITE = RGBColor(0xFF,0xFF,0xFF)
CREAM = RGBColor(0xF7,0xEE,0xCE); PAPER = RGBColor(0xF7,0xF7,0xF7)
GREENP= RGBColor(0xE9,0xEE,0xE1)
CG = "Century Gothic"; TNR = "Times New Roman"

prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=13.333,7.5; BLANK=prs.slide_layouts[6]; LM=0.75

def slide():
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb=PAPER; r.line.fill.background(); r.shadow.inherit=False
    sp=r._element; sp.getparent().remove(sp); s.shapes._spTree.insert(2,sp)
    try:
        p=s.shapes.add_picture(BGTEX,0,0,width=prs.slide_width,height=prs.slide_height)
        pe=p._element; pe.getparent().remove(pe); s.shapes._spTree.insert(3,pe)
    except Exception: pass
    return s

def txt(slide,l,t,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,line_spacing=1.0,space_after=0):
    tb=slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ('left','right','top','bottom'): setattr(tf,'margin_'+m,0)
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        if line_spacing: p.line_spacing=line_spacing
        p.space_before=0; p.space_after=Pt(space_after)
        for (text,st) in para:
            r=p.add_run(); r.text=text; f=r.font
            f.name=st.get('font',CG); f.size=Pt(st.get('size',18))
            f.bold=st.get('bold',False); f.italic=st.get('italic',False)
            f.color.rgb=st.get('color',DARK)
    return tb

def title(slide, runs_lines, size=24, y=0.55):
    """runs_lines: list of lines, each a list of (text,color) with claim styling."""
    paras=[]
    for line in runs_lines:
        paras.append([(t, dict(font=CG,size=size,color=c)) for (t,c) in line])
    txt(slide, LM, y, 12.0, 1.6, paras, line_spacing=1.1)

def citation(slide,text):
    txt(slide,LM,7.02,10,0.3,[[(text,dict(font=CG,size=10,color=FAINT))]])
def pagenum(slide,n):
    txt(slide,SW-LM-0.6,7.02,0.6,0.3,[[(str(n),dict(font=CG,size=10,color=FAINT))]],align=PP_ALIGN.RIGHT)

def rectshape(slide,l,t,w,h,fill,sharp=False,line=None,radius=0.1):
    shp=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE if sharp else MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(l),Inches(t),Inches(w),Inches(h))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(1.5)
    shp.shadow.inherit=False
    if not sharp:
        try: shp.adjustments[0]=radius
        except Exception: pass
    return shp

def box(slide,l,t,w,h,lines,fill=GREEN,tcolor=WHITE,size=15,bold=False,sharp=False,line=None,radius=0.12):
    rectshape(slide,l,t,w,h,fill,sharp=sharp,line=line,radius=radius)
    if isinstance(lines,str): lines=[lines]
    txt(slide,l+0.1,t,w-0.2,h,[[(ln,dict(font=CG,size=size,color=tcolor,bold=bold))] for ln in lines],
        align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.05)

def quotebox(slide,l,t,w,h,quote,tag,fill=GREEN,sharp=False,qcolor=WHITE,tagcolor=None):
    rectshape(slide,l,t,w,h,fill,sharp=sharp,radius=0.08)
    txt(slide,l+0.3,t+0.22,w-0.6,h-0.75,[[("“"+quote+"”",dict(font=TNR,size=16,italic=True,color=qcolor))]],
        anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.14)
    txt(slide,l+0.3,t+h-0.5,w-0.6,0.35,[[(tag,dict(font=CG,size=12,bold=True,color=tagcolor or qcolor))]])

def arrow(slide,x1,y1,x2,y2,color=GRAY,w=1.75,double=False):
    c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w); c.shadow.inherit=False
    ln=c.line._get_or_add_ln()
    for tg in (('a:headEnd','a:tailEnd') if double else ('a:tailEnd',)):
        ln.append(ln.makeelement(qn(tg),{'type':'triangle','w':'med','len':'med'}))

# ============================================= 1 TITLE
s=slide()
txt(s,1,1.05,11.33,0.4,[[("BCCE 2026",dict(font=CG,size=13,color=GREEND))]],align=PP_ALIGN.CENTER)
txt(s,1,2.35,11.33,2.0,[[("What chemistry faculty believe",dict(font=CG,size=38,color=DARK))],
    [("about doctoral education",dict(font=CG,size=38,color=DARK))]],align=PP_ALIGN.CENTER,line_spacing=1.1)
u=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(5.67),Inches(4.15),Inches(2.0),Inches(0.05))
u.fill.solid(); u.fill.fore_color.rgb=GREEN; u.line.fill.background(); u.shadow.inherit=False
txt(s,1,4.4,11.33,0.5,[[("A phenomenographic and thematic characterization",dict(font=CG,size=19,italic=True,color=GRAY))]],align=PP_ALIGN.CENTER)
txt(s,1,5.5,11.33,0.4,[[("Zoe Rahimi Pirkoohi",dict(font=CG,size=18,color=DARK))]],align=PP_ALIGN.CENTER)
txt(s,1,5.98,11.33,0.4,[[("Jordan Harshman  ·  University of Iowa",dict(font=CG,size=15,color=GRAY))]],align=PP_ALIGN.CENTER)
txt(s,1,6.95,11.33,0.35,[[("Supported by NSF CAREER #2142873 & #2602955  ·  Jordan Harshman, PI",dict(font=CG,size=10,color=FAINT))]],align=PP_ALIGN.CENTER)

# ============================================= 2 STAKES
s=slide()
title(s,[[("Doctoral education trains the chemists who staff",DARK)],
         [("the discipline — and the enterprise is straining.",DARK)]])
for i,(big,small) in enumerate([("~2,900","chemistry PhDs each year"),("~196","U.S. doctoral programs"),
        ("unchanged","the same structure for decades")]):
    x=LM+i*4.0
    box(s,x,2.75,3.7,1.35,[big,small],fill=GREEN,size=15)
txt(s,LM,4.65,12.0,1.2,
    [[("Faculty name four strains: ",dict(font=CG,size=16,color=DARK,bold=True)),
      ("balancing competing responsibilities, no standard assessment of outcomes, "
       "uneven implementation, and inconsistent, largely untrained mentorship.",dict(font=CG,size=16,color=DARK))]],line_spacing=1.3)
citation(s,"Donkor & Harshman, 2023  ·  Collini et al., 2025"); pagenum(s,2)

# ============================================= 3 REFORM FAILS -> BELIEFS
s=slide()
title(s,[[("Reform changes the ",DARK),("structures",GREEND),(" of the PhD, not what happens",DARK)],
         [("inside — because ",DARK),("beliefs absorb it",GREEND),(" without yielding.",DARK)]])
box(s,LM,2.95,3.35,1.15,["A reform arrives"],fill=GREEN,size=15)
arrow(s,4.15,3.52,4.75,3.52,color=GRAY,w=2)
box(s,4.8,2.95,3.35,1.15,["Structures change"],fill=GREEN,size=15)
arrow(s,8.2,3.52,8.8,3.52,color=GRAY,w=2)
box(s,8.85,2.95,3.73,1.15,["Practice unchanged"],fill=DARK,sharp=True,size=15)
box(s,2.4,4.7,8.5,1.0,["Beliefs are self-perpetuating — they can absorb a reform","without yielding to it. So change must start with them."],fill=CREAM,tcolor=DARK,size=16)
citation(s,"Teacher-Centered Systemic Reform · Woodbury & Gess-Newsome, 2002"); pagenum(s,3)

# ============================================= 4 THE GAP
s=slide()
title(s,[[("We know what faculty believe about ",DARK),("teaching a class",GREEND),(".",DARK)],
         [("No one has characterized what they believe about ",DARK),("advising a PhD",GREEND),(".",DARK)]])
box(s,LM,3.0,5.6,1.6,["TEACHING A CLASS","","Beliefs well characterized —","and they predict practice"],fill=GREEN,size=15)
box(s,LM+6.2,3.0,5.6,1.6,["ADVISING A PhD","","A different practice —","never characterized"],fill=DARK,sharp=True,size=15)
txt(s,LM,5.2,12.0,0.8,[[("The closest prior study set beliefs aside — a proper look ",dict(font=CG,size=15,color=DARK)),
    ("“warranted its own study.”",dict(font=TNR,size=16,italic=True,color=GREEND)),
    ("  This is that study.",dict(font=CG,size=15,color=DARK,bold=True))]],line_spacing=1.25)
citation(s,"Kagan, 1992  ·  Donkor et al., 2024"); pagenum(s,4)

# ============================================= 5 RQ
s=slide()
txt(s,1,1.6,11.33,0.4,[[("RESEARCH QUESTION",dict(font=CG,size=13,color=GREEND))]],align=PP_ALIGN.CENTER)
txt(s,1,2.6,11.33,1.6,[[("What do chemistry faculty believe",dict(font=CG,size=32,color=DARK))],
    [("about doctoral education?",dict(font=CG,size=32,color=DARK))]],align=PP_ALIGN.CENTER,line_spacing=1.12)
u=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(6.17),Inches(4.55),Inches(1.0),Inches(0.05))
u.fill.solid(); u.fill.fore_color.rgb=GREEN; u.line.fill.background(); u.shadow.inherit=False
txt(s,1,4.9,11.33,0.5,[[("what it’s for   ·   how they learn   ·   who develops which skill   ·   the structure they work in",dict(font=CG,size=14.5,color=GRAY))]],align=PP_ALIGN.CENTER)
pagenum(s,5)

# ============================================= 6 WHAT IS A BELIEF
s=slide()
title(s,[[("First, what counts as a ",DARK),("belief",GREEND),("? A proposition a person holds",DARK)],
         [("true and acts on — not a fact, and not easily changed.",DARK)]])
for i,(term,gloss) in enumerate([("Propositional","held to be true"),("Evaluative","guides decisions"),
        ("Personal","not group consensus"),("Resistant","survives counter-evidence")]):
    x=LM+i*2.95
    box(s,x,3.05,2.6,1.15,[term,"",gloss],fill=GREEN,size=14.5)
    if i<3: arrow(s,x+2.62,3.62,x+2.93,3.62,color=GRAY,w=1.5)
box(s,5.4,4.65,2.5,0.7,["= a belief"],fill=DARK,sharp=True,size=16,bold=True)
txt(s,LM,5.65,12,0.5,[[("We inferred beliefs from what faculty said and why — a knowledge claim answers to evidence; a belief is held on conviction.",dict(font=CG,size=13,italic=True,color=GRAY))]],align=PP_ALIGN.CENTER)
citation(s,"Rokeach, 1968 · Green, 1971 · Nespor, 1987 · Pajares, 1992"); pagenum(s,6)

# ============================================= 7 METHOD
s=slide()
title(s,[[("We built those beliefs from twenty faculty’s",DARK)],
         [("own words, coded to saturation.",DARK)]])
for i,(a,b) in enumerate([("Twenty faculty","U.S. doctoral programs, still advising"),
        ("Interviews + card sort","~60 min, in their own words"),
        ("Phenomenography +","thematic analysis")]):
    box(s,LM+i*4.0,2.8,3.7,1.2,[a,"",b],fill=GREEN,size=14)
for i,(big,small) in enumerate([("990","sub-codes"),("15","codes"),("4","themes")]):
    x=2.15+i*2.5
    box(s,x,4.65,2.0,0.9,[big+"  "+small],fill=None,tcolor=DARK,line=GREEN,size=16)
    if i<2: arrow(s,x+2.02,5.1,x+2.48,5.1,color=GREEN,w=1.75)
txt(s,LM,5.75,12,0.4,[[("No new codes appeared after interview 7 — the result held from 10 participants to all 20.",dict(font=CG,size=13,italic=True,color=GRAY))]],align=PP_ALIGN.CENTER)
citation(s,"Braun & Clarke, 2006 · Åkerlind, 2005"); pagenum(s,7)

# ============================================= 8 THE TWIST
s=slide()
title(s,[[("Here is the finding: faculty ",DARK),("agree",GREEND),(" on the propositions —",DARK)],
         [("and ",DARK),("divide",GREEND),(" on exactly where each boundary falls.",DARK)]])
box(s,LM,3.05,5.6,1.55,["AGREE","on four shared propositions"],fill=GREEN,size=17)
box(s,LM+6.2,3.05,5.6,1.55,["DIVIDE","on where the boundaries fall"],fill=DARK,sharp=True,size=17)
txt(s,LM,5.1,12.0,1.0,[[("“What differs is not which beliefs faculty hold, but where they draw the lines those beliefs require.”",dict(font=TNR,size=18,italic=True,color=GREEND))]],align=PP_ALIGN.CENTER,line_spacing=1.2)
pagenum(s,8)

# ============================================= 9 BELIEF 1
s=slide()
txt(s,LM,0.5,12,0.3,[[("BELIEF 1  ·  WHAT THE DOCTORATE IS FOR",dict(font=CG,size=10.5,color=GREEND))]])
title(s,[[("The PhD builds a ",DARK),("self-directed scientist",GREEND),(" —",DARK)],
         [("in service of the student’s own goals.",DARK)]],y=0.95)
quotebox(s,LM,2.85,7.3,1.55,"When they start bringing ideas forward without necessarily being prompted.","P2 — on what independence looks like",fill=GREEN)
box(s,8.35,2.85,4.23,1.55,["“It’s not my career.","It’s not my goals.”","— P2, on whose goals lead"],fill=GREEND,size=14)
txt(s,LM,4.75,8,0.3,[[("WHERE THEY DIVIDE",dict(font=CG,size=10.5,color=GREEND))]])
arrow(s,2.4,5.35,10.9,5.35,color=GRAY,double=True)
txt(s,2.3,5.45,4.6,0.4,[[("answering existing questions",dict(font=CG,size=13.5,color=DARK))]])
txt(s,6.4,5.45,4.5,0.4,[[("asking new questions",dict(font=CG,size=13.5,color=DARK))]],align=PP_ALIGN.RIGHT)
txt(s,LM,5.9,12,0.4,[[("The realistic goal of the degree, versus the higher bar. Independence is the shared claim; its reach is the split.",dict(font=CG,size=12,italic=True,color=GRAY))]])
pagenum(s,9)

# ============================================= 10 BELIEF 2 (contrast)
s=slide()
txt(s,LM,0.5,12,0.3,[[("BELIEF 2  ·  HOW STUDENTS LEARN",dict(font=CG,size=10.5,color=GREEND))]])
title(s,[[("Students learn by ",DARK),("doing scaffolded research",GREEND),(" — but the",DARK)],
         [("same lonely start taught two faculty opposite lessons.",DARK)]],y=0.95)
quotebox(s,LM,2.9,5.9,2.0,"Finding the answers on my own was not the most efficient — but it made me a much better scientist.","P13 — credited the struggle",fill=GREEN)
quotebox(s,LM+6.5,2.9,5.83,2.0,"Moving toward independence is really done in stages. Not something I tried to model at all.","P7 — drew the opposite lesson",fill=DARK,sharp=True)
txt(s,LM,5.25,12,0.4,[[("Both were left alone as students. ",dict(font=CG,size=14,color=DARK,bold=True)),
    ("The shared belief is learning-by-doing; the divide is how much unsupported struggle helps.",dict(font=CG,size=14,color=DARK))]],line_spacing=1.25)
citation(s,"P5: “I do. We do. You do.”  ·  P12: “Can you teach someone else to do it?”"); pagenum(s,10)

# ============================================= 11 BELIEF 3 (contrast)
s=slide()
txt(s,LM,0.5,12,0.3,[[("BELIEF 3  ·  WHO DEVELOPS WHICH SKILL — THE LARGEST THEME",dict(font=CG,size=10.5,color=GREEND))]])
title(s,[[("Development is a ",DARK),("shared, two-way job",GREEND),(" — yet the same skill",DARK)],
         [("is one advisor’s to teach and another’s for the student to build.",DARK)]],y=0.95)
box(s,LM,2.95,5.9,1.75,["“Creativity is the faculty’s to teach.”","","— P5 · P17"],fill=GREEN,size=15)
box(s,LM+6.5,2.95,5.83,1.75,["“Creativity is the student’s to build.”","","— P9 · P11"],fill=DARK,sharp=True,size=15)
txt(s,LM,5.05,12,0.4,[[("Every advisor draws a line between what they must teach and what the student must build — ",dict(font=CG,size=14,color=DARK)),
    ("and they draw it in different places.",dict(font=CG,size=14,color=DARK,bold=True))]],line_spacing=1.25)
citation(s,"P6: “I go from the leader to the follower…”  ·  P13: “I cannot teach you to be excited about the science.”"); pagenum(s,11)

# ============================================= 12 BELIEF 4
s=slide()
txt(s,LM,0.5,12,0.3,[[("BELIEF 4  ·  THE STRUCTURAL CONDITIONS",dict(font=CG,size=10.5,color=GREEND))]])
title(s,[[("Structure shapes advising — and the ",DARK),("gaps",GREEND),(" it leaves",DARK)],
         [("are the program’s to fix, not the student’s.",DARK)]],y=0.95)
quotebox(s,LM,2.9,7.3,1.5,"Those professional skills often get kind of forgotten.","P2 — on the program’s gap",fill=GREEN)
box(s,8.35,2.9,4.23,1.5,["Fourteen of twenty","already name the gaps —","and can fix them."],fill=GREEND,size=14)
txt(s,LM,4.7,8,0.3,[[("WHERE THEY DIVIDE",dict(font=CG,size=10.5,color=GREEND))]])
arrow(s,2.4,5.3,10.9,5.3,color=GRAY,double=True)
txt(s,2.3,5.4,4.6,0.4,[[("the evolved system works",dict(font=CG,size=13.5,color=DARK))]])
txt(s,6.4,5.4,4.5,0.4,[[("curricula & funding need change",dict(font=CG,size=13.5,color=DARK))]],align=PP_ALIGN.RIGHT)
txt(s,LM,5.85,12,0.4,[[("On structure, many faculty are already critics rather than obstacles — they disagree on which gaps are real.",dict(font=CG,size=12,italic=True,color=GRAY))]])
pagenum(s,12)

# ============================================= 13 THE MAP
s=slide()
s.shapes.add_picture(MAP,Inches(0.6),Inches(0.9),height=Inches(6.1))
RX=7.15
title([s][0], [[("Agreement on every theme —",DARK)]], size=23, y=1.05) if False else None
txt(s,RX,1.05,5.6,1.2,[[("Agreement on every theme.",dict(font=CG,size=23,color=DARK))],
    [("A ",dict(font=CG,size=23,color=DARK)),("fault line",dict(font=CG,size=23,color=GREEND)),(" under every one.",dict(font=CG,size=23,color=DARK))]],line_spacing=1.15)
txt(s,RX,3.15,5.6,2.4,[[("Fifteen codes. Fifteen axes of division. ",dict(font=CG,size=15,color=DARK,bold=True)),
    ("Every faculty member lands somewhere on each — the agreement is real, and so is the disagreement one level down.",dict(font=CG,size=15,color=GRAY))]],line_spacing=1.3)
pagenum(s,13)

# ============================================= 14 IMPLICATIONS
s=slide()
title(s,[[("This gives reform its real map: shared ground to",DARK)],
         [("build on, and where “aligned” colleagues diverge.",DARK)]])
box(s,LM,2.9,5.6,1.55,["SHARED PROPOSITIONS","","Common ground — argue from where","faculty already stand."],fill=GREEN,size=14.5)
box(s,LM+6.2,2.9,5.6,1.55,["BOUNDARY DISAGREEMENTS","","The real work — surface them, don’t","assume they’re settled."],fill=DARK,sharp=True,size=14.5)
box(s,2.4,4.75,8.5,1.05,["And convictions stay stable while practice is refined —","so support the implementation, not just the belief."],fill=CREAM,tcolor=DARK,size=15)
citation(s,"13 of 20: stable convictions, refined implementation  ·  Popova et al., 2021"); pagenum(s,14)

# ============================================= 15 CONCLUSION
s=slide()
txt(s,1,1.5,11.33,0.4,[[("IN ONE LINE",dict(font=CG,size=13,color=GREEND))]],align=PP_ALIGN.CENTER)
txt(s,1,2.5,11.33,1.8,[[("Four shared convictions about the PhD —",dict(font=CG,size=28,color=DARK))],
    [("and the lines faculty quietly draw differently.",dict(font=CG,size=28,color=DARK))]],align=PP_ALIGN.CENTER,line_spacing=1.15)
u=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(6.17),Inches(4.5),Inches(1.0),Inches(0.05))
u.fill.solid(); u.fill.fore_color.rgb=GREEN; u.line.fill.background(); u.shadow.inherit=False
txt(s,1,4.85,11.33,0.5,[[("Thank you.",dict(font=CG,size=22,color=DARK))]],align=PP_ALIGN.CENTER)
txt(s,1,5.55,11.33,0.4,[[("Zoe Rahimi Pirkoohi  ·  Jordan Harshman  ·  University of Iowa",dict(font=CG,size=14,color=GRAY))]],align=PP_ALIGN.CENTER)
txt(s,1,6.9,11.33,0.35,[[("Supported by NSF CAREER #2142873 & #2602955  ·  Jordan Harshman, PI",dict(font=CG,size=10,color=FAINT))]],align=PP_ALIGN.CENTER)

out=os.path.join(HERE,"BCCE2026_faculty_beliefs.pptx"); prs.save(out)
print("saved",out,len(prs.slides._sldIdLst),"slides")
