#!/usr/bin/env python3
"""Minimal pptx -> PNG renderer for visual QA (parses the real pptx)."""
import sys, io, os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SCALE = 110  # px per inch
def px(emu): return int(Emu(emu).inches * SCALE)

SERIF = {
 (0,0): "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
 (1,0): "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
 (0,1): "/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf",
 (1,1): "/usr/share/fonts/truetype/liberation/LiberationSerif-BoldItalic.ttf",
}
SANS = {
 (0,0): "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
 (1,0): "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
 (0,1): "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
 (1,1): "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf",
}
_fc = {}
def font(name, size_pt, bold, italic):
    fam = SERIF if (name or "").lower().startswith("georgia") else SANS
    path = fam[(1 if bold else 0, 1 if italic else 0)]
    key = (path, int(size_pt*SCALE/72))
    if key not in _fc:
        _fc[key] = ImageFont.truetype(path, key[1])
    return _fc[key]

def rgb(c):
    try: return (c[0], c[1], c[2])
    except Exception: return None

def shape_fill(shape):
    try:
        if shape.fill.type is not None and shape.fill.type == 1:  # solid
            return (shape.fill.fore_color.rgb[0], shape.fill.fore_color.rgb[1], shape.fill.fore_color.rgb[2])
    except Exception: pass
    return None

def shape_line(shape):
    try:
        if shape.line.color and shape.line.color.type is not None:
            c = shape.line.color.rgb
            w = shape.line.width
            return (c[0],c[1],c[2]), max(1, int(Emu(w).inches*SCALE)) if w else 1
    except Exception: pass
    return None

def draw_text(draw, shape):
    tf = shape.text_frame
    L = px(shape.left); T = px(shape.top); W = px(shape.width); H = px(shape.height)
    ml = px(tf.margin_left or 0); mr = px(tf.margin_right or 0)
    mt = px(tf.margin_top or 0); mb = px(tf.margin_bottom or 0)
    x0 = L+ml; y0 = T+mt; boxw = W-ml-mr
    anchor = str(tf.vertical_anchor)
    # build lines
    lines = []  # list of (tokens) where token=(text,font,color,w,h)
    for para in tf.paragraphs:
        align = str(para.alignment)
        ls = para.line_spacing or 1.0
        toks = []
        for run in para.runs:
            fnt = font(run.font.name, run.font.size.pt if run.font.size else 18,
                       bool(run.font.bold), bool(run.font.italic))
            try: col = rgb(run.font.color.rgb) or (30,30,30)
            except Exception: col = (30,30,30)
            words = run.text.split(" ")
            for i,w in enumerate(words):
                t = w + (" " if i < len(words)-1 else "")
                if t == "": continue
                bb = draw.textbbox((0,0), t, font=fnt)
                toks.append([t, fnt, col, bb[2]-bb[0], bb[3]-bb[1], fnt.size])
        # wrap
        cur=[]; curw=0
        wrapped=[]
        for tk in toks:
            if curw+tk[3] > boxw and cur:
                wrapped.append(cur); cur=[tk]; curw=tk[3]
            else:
                cur.append(tk); curw+=tk[3]
        if cur: wrapped.append(cur)
        if not wrapped: wrapped=[[]]
        for wl in wrapped:
            lines.append((wl, align, ls))
    # total height
    def lineh(wl, ls):
        h = max([t[5] for t in wl], default=int(18*SCALE/72))
        return h*1.0*ls + h*0.35
    total = sum(lineh(wl,ls) for wl,_,ls in lines)
    if "BOTTOM" in anchor: y = y0 + (H-mt-mb) - total
    elif "MIDDLE" in anchor: y = y0 + ((H-mt-mb)-total)/2
    else: y = y0
    for wl, align, ls in lines:
        lw = sum(t[3] for t in wl)
        if "CENTER" in align: x = x0 + (boxw-lw)/2
        elif "RIGHT" in align: x = x0 + (boxw-lw)
        else: x = x0
        h = max([t[5] for t in wl], default=int(18*SCALE/72))
        for t in wl:
            draw.text((x, y+ (h-t[4])*0.2), t[0], font=t[1], fill=t[2])
            x += t[3]
        y += lineh(wl, ls)

def render(pptx_path, outdir):
    prs = Presentation(pptx_path)
    W = px(prs.slide_width); H = px(prs.slide_height)
    paths=[]
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W,H), (252,251,247))
        d = ImageDraw.Draw(img, "RGBA")
        for shape in slide.shapes:
            tag = shape._element.tag
            try:
                L=px(shape.left); T=px(shape.top); Wd=px(shape.width); Hd=px(shape.height)
            except Exception:
                continue
            if tag.endswith('}cxnSp'):
                lc = shape_line(shape)
                if lc: d.line([(L,T),(L+Wd,T+Hd)], fill=lc[0], width=lc[1])
                continue
            if shape.shape_type == 13 or tag.endswith('}pic'):  # picture
                try:
                    blob = shape.image.blob
                    im2 = Image.open(io.BytesIO(blob)).convert("RGBA")
                    im2 = im2.resize((max(1,Wd), max(1,Hd)))
                    img.paste(im2, (L,T), im2)
                except Exception as e:
                    d.rectangle([L,T,L+Wd,T+Hd], outline=(200,120,120), width=2)
                continue
            # autoshape
            fill = shape_fill(shape); ln = shape_line(shape)
            st = None
            try: st = shape.shape_type
            except Exception: pass
            name = ""
            try: name = shape._element.find('.//{*}prstGeom').get('prst')
            except Exception: pass
            box=[L,T,L+Wd,T+Hd]
            if name == 'ellipse':
                if fill: d.ellipse(box, fill=fill)
                if ln: d.ellipse(box, outline=ln[0], width=ln[1])
            elif name == 'roundRect':
                rad = min(Wd,Hd)*0.18
                if fill: d.rounded_rectangle(box, radius=rad, fill=fill)
                if ln: d.rounded_rectangle(box, radius=rad, outline=ln[0], width=ln[1])
            elif name == 'chevron':
                if fill:
                    d.polygon([(L,T),(L+Wd*0.6,T),(L+Wd,T+Hd/2),(L+Wd*0.6,T+Hd),(L,T+Hd)], fill=fill)
            else:  # rectangle etc
                if fill: d.rectangle(box, fill=fill)
                if ln: d.rectangle(box, outline=ln[0], width=ln[1])
            # text
            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_text(d, shape)
        p = os.path.join(outdir, f"slide_{idx:02d}.png")
        img.save(p); paths.append(p)
    # contact sheet
    cols=4; rows=(len(paths)+cols-1)//cols
    tw=W//3; th=H//3
    sheet=Image.new("RGB",(tw*cols, th*rows),(230,230,228))
    for i,p in enumerate(paths):
        im=Image.open(p).resize((tw,th)); sheet.paste(im,((i%cols)*tw,(i//cols)*th))
    sp=os.path.join(outdir,"contact_sheet.png"); sheet.save(sp)
    print("rendered", len(paths), "slides ->", sp)
    return paths

if __name__ == "__main__":
    render(sys.argv[1], os.path.dirname(os.path.abspath(sys.argv[1])))
